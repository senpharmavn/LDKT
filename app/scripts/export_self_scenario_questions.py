import csv
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Evaluation Round 2 Report.xlsx"
SELF_OUTPUT = ROOT / "app" / "docs" / "evaluation_self_scenario_questions.csv"
PEER_OUTPUT = ROOT / "app" / "docs" / "evaluation_peer_scenario_questions.csv"


def main() -> None:
    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    sheet = workbook["Main_Key"]

    self_rows = []
    peer_rows = []
    dimension = ""
    for row_number in range(6, 36):
        raw_dimension = sheet.cell(row_number, 1).value
        if raw_dimension:
            dimension = str(raw_dimension).replace("\n", " ").strip()

        order = sheet.cell(row_number, 2).value
        scenario = sheet.cell(row_number, 3).value
        reaction = sheet.cell(row_number, 4).value
        creative = sheet.cell(row_number, 5).value
        peer_scenario = sheet.cell(row_number, 7).value
        peer_reaction = sheet.cell(row_number, 8).value
        peer_creative = sheet.cell(row_number, 9).value
        if not order or not scenario:
            continue

        self_rows.append(
            [
                f"SELF-SCENARIO-Q{int(order):02d}",
                "SELF-COMMON",
                int(order),
                str(scenario).strip(),
                "single_choice",
                f"Phản ứng::{str(reaction).strip()}|Sáng tạo::{str(creative).strip()}",
                "TRUE",
                dimension,
                "round2",
            ]
        )

        if peer_scenario and peer_reaction and peer_creative:
            peer_rows.append(
                [
                    f"PEER-SCENARIO-Q{int(order):02d}",
                    "PEER-COMMON",
                    int(order),
                    str(peer_scenario).strip(),
                    "single_choice",
                    f"Phản ứng::{str(peer_reaction).strip()}|Sáng tạo::{str(peer_creative).strip()}",
                    "TRUE",
                    dimension,
                    "round2",
                ]
            )

    safe_write_questions_csv(SELF_OUTPUT, self_rows)
    safe_write_questions_csv(PEER_OUTPUT, peer_rows)
    print(f"Exported {len(self_rows)} self questions to {SELF_OUTPUT}")
    print(f"Exported {len(peer_rows)} peer questions to {PEER_OUTPUT}")


def write_questions_csv(output: Path, rows: list[list[object]]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "QuestionID",
                "TemplateID",
                "Order",
                "QuestionText",
                "QuestionType",
                "Options",
                "Required",
                "Dimension",
                "Version",
            ]
        )
        writer.writerows(rows)


def safe_write_questions_csv(output: Path, rows: list[list[object]]) -> None:
    try:
        write_questions_csv(output, rows)
    except PermissionError:
        print(f"Skipped locked file: {output}")


if __name__ == "__main__":
    main()
